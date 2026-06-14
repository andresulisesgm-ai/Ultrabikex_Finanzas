import sqlite3
import openpyxl
import os

db_path = 'data/ultrax.db'
excel_path = 'EEFF ULTRAX 2026.xlsx'

print("="*80)
print("SECCIÓN 1 - Tabla mapping (V1):")
print("="*80)
try:
    conn = sqlite3.connect(db_path)
    c = conn.cursor()
    c.execute("SELECT * FROM mapping WHERE odoo_code='6.01.02.03.002'")
    rows = c.fetchall()
    print("Column names: odoo_code, odoo_name, partida, sign, income_type")
    for r in rows:
        print(r)
except Exception as e:
    print("Error in Sección 1:", e)

print("\n" + "="*80)
print("SECCIÓN 2 - Esquema de mapping_groups_v2:")
print("="*80)
try:
    c.execute("PRAGMA table_info(mapping_groups_v2)")
    columns = c.fetchall()
    for col in columns:
        print(f"Col {col[0]}: {col[1]} ({col[2]})")
except Exception as e:
    print("Error in Sección 2:", e)

print("\n" + "="*80)
print("SECCIÓN 3 - Ejemplo de cuenta similar ya mapeada en mapping_groups_v2:")
print("="*80)
try:
    # Since mapping_groups_v2 uses group_name, we query by group_name.
    # We will also try odoo_code mapped to odoo_name in mapping.
    c.execute("""
        SELECT * FROM mapping_groups_v2 
        WHERE group_name LIKE '%personal%' OR group_name LIKE '%administra%' 
        LIMIT 5
    """)
    rows = c.fetchall()
    print("Querying by group_name in mapping_groups_v2:")
    for r in rows:
        print(r)
        
    print("\nQuery joining mapping_groups_v2 with mapping to search by odoo_name:")
    c.execute("""
        SELECT m2.*, m1.odoo_name
        FROM mapping_groups_v2 m2
        JOIN mapping m1 ON m2.odoo_code = m1.odoo_code
        WHERE m1.odoo_name LIKE '%personal%' OR m1.odoo_name LIKE '%administra%'
        LIMIT 5
    """)
    rows_join = c.fetchall()
    for r in rows_join:
        print(r)
except Exception as e:
    print("Error in Sección 3:", e)

print("\n" + "="*80)
print("SECCIÓN 4 - Contexto en N EERR alrededor de la fila A132:")
print("="*80)
try:
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['N EERR']
    for r_idx in range(110, 141):
        cell_a = ws.cell(row=r_idx, column=1)
        cell_b = ws.cell(row=r_idx, column=2)
        val_a = cell_a.value
        val_b = cell_b.value
        if val_a is not None or val_b is not None:
            print(f"Row {r_idx}: A={repr(val_a)}, B={repr(val_b)}")
except Exception as e:
    print("Error in Sección 4:", e)

if conn:
    conn.close()
