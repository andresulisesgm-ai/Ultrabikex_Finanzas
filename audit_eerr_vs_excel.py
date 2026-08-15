"""
audit_eerr_vs_excel.py
Compara EERR normal (BCV) del sistema contra el Excel oficial de Yocelin,
las 6 unidades de negocio, ENE-JUN 2026. Imprime SOLO discrepancias > $0.02.
"""
import sys
import os
import json
import unicodedata
from openpyxl import load_workbook

# Configurar path e importar app de Flask para autenticación en scripts
sys.path.insert(0, r"C:\Users\andre\Desktop\ultrax_app")
sys.stderr = open(os.devnull, 'w')
from app import app

EXCEL_PATH = r"C:\Users\andre\Desktop\ultrax_app\2 EEFF ULTRAX COMPLETO 2026.xlsx"
BASE_URL = "http://localhost:5000"
TOLERANCIA = 0.02

# Mapeo unidad -> nombre de hoja Excel (hoja SIN prefijo "N ", resumen/oficial)
UNIDADES = {
    "Rodeo": "EERR RODEO",
    "PiedeMonte": "EERR PIEDEM",
    "Terracota": "EERR TERRA",
    "Ucafe": "EERR UCAFE",
    "Barinas": "EERR BARINAS",
    "Naranjos": "EERR LOS NA",
}

# Meses con datos reales cargados (ENE-JUN). JUL en adelante no tiene datos aún.
MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN"]

# Columna mensual (ejecutado) y columna acumulado (ACUM EJEC) por mes,
# según encabezado real fila 3 de "EERR RODEO". El script relee la fila 3
# de CADA hoja antes de usar estas columnas, y aborta esa unidad con aviso
# si el encabezado no coincide con este patrón (nunca asume a ciegas).
COL_MENSUAL = {"ENE": "E", "FEB": "H", "MAR": "O", "ABR": "Y", "MAY": "AF", "JUN": "AM"}
COL_ACUM = {"FEB": "L", "MAR": "S", "ABR": "AC", "MAY": "AJ", "JUN": "AQ"}  # ENE no tiene acumulado propio


def normalizar(texto):
    if texto is None:
        return ""
    texto = texto.strip()
    texto = texto.lower()
    texto = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode('ascii')  # quita acentos
    texto = ' '.join(texto.split())  # colapsa espacios múltiples
    return texto


def verificar_encabezado(ws, hoja_nombre):
    """Verifica que la fila 3 tenga el patrón esperado en las columnas clave antes de leer nada."""
    esperado = {
        "E3": "ENE", "H3": "FEB", "O3": "MAR", "Y3": "ABR", "AF3": "MAY", "AM3": "JUN",
        "L3": "ACUM EJEC", "S3": "ACUM EJEC", "AC3": "ACUM EJEC", "AJ3": "ACUM EJEC", "AQ3": "ACUM EJEC",
    }
    errores = []
    for celda, esperado_val in esperado.items():
        real = ws[celda].value
        if real != esperado_val:
            errores.append(f"{celda}: esperado '{esperado_val}', encontrado '{real}'")
    if errores:
        print(f"[ABORTADO] Hoja '{hoja_nombre}' no coincide con el patrón de encabezado esperado:")
        for e in errores:
            print(f"    {e}")
        return False
    return True


def cargar_excel_valores(ws):
    """Devuelve dict: {partida_nombre: {mes: {'mensual': v, 'acum': v}}}"""
    datos = {}
    for row in ws.iter_rows(min_row=4, max_col=1):
        celda = row[0]
        nombre = celda.value
        if not nombre:
            continue
        fila = celda.row
        datos[nombre] = {}
        for mes in MESES:
            col_m = COL_MENSUAL[mes]
            v_mensual = ws[f"{col_m}{fila}"].value
            entry = {"mensual": v_mensual}
            if mes in COL_ACUM:
                col_a = COL_ACUM[mes]
                entry["acum"] = ws[f"{col_a}{fila}"].value
            datos[nombre][mes] = entry
    return datos


def cargar_sistema_valores(unit, empresa_id=2, year=2026):
    """Devuelve dict: {partida_nombre: {mes: {'mensual': v, 'acum': v}}}"""
    client = app.test_client()
    with client.session_transaction() as sess:
        sess['username'] = 'Yoce'
        sess['role'] = 'admin'

    url = f"/api/eerr/completo?year={year}&unit={unit}&empresa_id={empresa_id}"
    resp = client.get(url)
    data = resp.get_json()
    datos = {}
    for row in data.get("rows", []):
        partida = row.get("partida")
        if partida in (None, "ESTADO DE RESULTADOS", "PARTIDAS"):
            continue
        datos[partida] = {}
        for mes_obj in row.get("meses", []):
            mes = mes_obj.get("month")
            if mes not in MESES:
                continue
            ejecutado = mes_obj.get("ejecutado", {}).get("valor")
            acum = mes_obj.get("acum_ejecutado", {}).get("valor") if "acum_ejecutado" in mes_obj else None
            datos[partida][mes] = {"mensual": ejecutado, "acum": acum}
    return datos


def comparar(unit, sistema, excel):
    discrepancias = []
    # Mapa normalizado -> nombre original en Excel
    excel_norm_map = {normalizar(k): k for k in excel.keys()}

    for partida, meses_sistema in sistema.items():
        partida_norm = normalizar(partida)
        if partida_norm not in excel_norm_map:
            continue  # no existe equivalente normalizado en Excel

        excel_key = excel_norm_map[partida_norm]
        meses_excel = excel.get(excel_key, {})

        for mes in MESES:
            s = meses_sistema.get(mes, {})
            e = meses_excel.get(mes, {})

            v_s_mensual = s.get("mensual")
            v_e_mensual = e.get("mensual")
            if v_s_mensual is not None and v_e_mensual is not None:
                try:
                    diff = abs(float(v_s_mensual) - float(v_e_mensual))
                    if diff > TOLERANCIA:
                        discrepancias.append((unit, partida, mes, "mensual", v_s_mensual, v_e_mensual, diff))
                except (TypeError, ValueError):
                    pass

            if mes in COL_ACUM:
                v_s_acum = s.get("acum")
                v_e_acum = e.get("acum")
                if v_s_acum is not None and v_e_acum is not None:
                    try:
                        diff = abs(float(v_s_acum) - float(v_e_acum))
                        if diff > TOLERANCIA:
                            discrepancias.append((unit, partida, mes, "acumulado", v_s_acum, v_e_acum, diff))
                    except (TypeError, ValueError):
                        pass
    return discrepancias


def main():
    wb = load_workbook(EXCEL_PATH, data_only=True)
    todas_discrepancias = []
    partidas_no_encontradas = []

    for unit, hoja_nombre in UNIDADES.items():
        print(f"\n=== Procesando {unit} (hoja: {hoja_nombre}) ===")
        if hoja_nombre not in wb.sheetnames:
            print(f"[ABORTADO] Hoja '{hoja_nombre}' no existe en el Excel.")
            continue
        ws = wb[hoja_nombre]
        if not verificar_encabezado(ws, hoja_nombre):
            continue

        excel_datos = cargar_excel_valores(ws)
        sistema_datos = cargar_sistema_valores(unit)

        excel_norm_set = {normalizar(k) for k in excel_datos.keys()}

        for partida in sistema_datos:
            if normalizar(partida) not in excel_norm_set:
                partidas_no_encontradas.append((unit, partida))

        discrepancias = comparar(unit, sistema_datos, excel_datos)
        todas_discrepancias.extend(discrepancias)

    print("\n\n========== RESULTADO FINAL ==========")
    print(f"\nTotal discrepancias (> ${TOLERANCIA}): {len(todas_discrepancias)}\n")
    for d in todas_discrepancias:
        unit, partida, mes, tipo, v_sis, v_exc, diff = d
        print(f"[{unit}] {partida} | {mes} ({tipo}) | Sistema: {v_sis} | Excel: {v_exc} | Diff: {diff:.2f}")

    if partidas_no_encontradas:
        print(f"\n\nPartidas del sistema SIN fila equivalente en Excel tras normalizar (revisar nombre, no es diferencia numérica): {len(partidas_no_encontradas)}")
        for unit, partida in partidas_no_encontradas:
            print(f"  [{unit}] '{partida}'")


if __name__ == "__main__":
    main()
