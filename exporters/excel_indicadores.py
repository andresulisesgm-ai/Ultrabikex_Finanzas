class IndicadoresExporter:
    def __init__(self, year):
        self.year = year

    def generate(self):
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from engine import compute_indicadores_v2
        from db import get_db
        import os, tempfile

        db = get_db()
        indicadores = compute_indicadores_v2(db, self.year)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'INDICADORES'
        ws.sheet_view.showGridLines = False

        HDR_FILL  = PatternFill('solid', start_color='1F3864')
        SUB_FILL  = PatternFill('solid', start_color='D6DCE4')
        WHITE_FONT= Font(name='Arial', bold=True, color='FFFFFF', size=10)
        DARK_FONT = Font(name='Arial', bold=True, color='000000', size=10)
        NORM_FONT = Font(name='Arial', size=9)
        BOLD_FONT = Font(name='Arial', bold=True, size=9)
        NUM_FMT   = '#,##0.00;(#,##0.00);"-"'
        PCT_FMT   = '0.0%;(0.0%);"-"'
        RATIO_FMT = '0.00;(0.00);"-"'
        thin      = Side(style='thin', color='D9D9D9')
        border    = Border(left=thin, right=thin, top=thin, bottom=thin)

        year_prev = str(int(self.year) - 1)
        QLABELS = {1: 'Q1 (Ene-Mar)', 2: 'Q2 (Abr-Jun)', 3: 'Q3 (Jul-Sep)', 4: 'Q4 (Oct-Dic)'}

        last_col = 12  # A..L
        ws.merge_cells(f'A1:{get_column_letter(last_col)}1')
        title = ws['A1']
        title.value = f'INDICADORES FINANCIEROS — CONSOLIDADO — {self.year}'
        title.font = Font(name='Arial', bold=True, color='FFFFFF', size=12)
        title.fill = HDR_FILL
        title.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 22

        headers = ['INDICADOR', year_prev]
        for q in [1, 2, 3, 4]:
            headers += [QLABELS[q], 'Vari Rel.']
        headers += [f'AÑO ACTUAL {self.year}', 'Vari Rel.']

        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=2, column=col, value=h)
            c.font = DARK_FONT
            c.fill = SUB_FILL
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = border

        ws.column_dimensions['A'].width = 42
        for col in range(2, last_col + 1):
            ws.column_dimensions[get_column_letter(col)].width = 13

        def fmt_valor(v, es_pct, es_ratio):
            if v is None:
                return None, NUM_FMT
            if es_pct:
                return v, PCT_FMT
            if es_ratio:
                return v, RATIO_FMT
            return v, NUM_FMT

        row = 3
        for ind in indicadores:
            nombre = ind['nombre']
            es_pct, es_ratio = ind['es_pct'], ind['es_ratio']

            c = ws.cell(row=row, column=1, value=nombre)
            c.font = BOLD_FONT
            c.border = border
            c.alignment = Alignment(horizontal='left', vertical='center')

            col = 2
            v, f = fmt_valor(ind['year_prev'], es_pct, es_ratio)
            c = ws.cell(row=row, column=col, value=v)
            c.number_format = f
            c.font = NORM_FONT
            c.border = border
            col += 1

            for t in ind['trimestres']:
                v, f = fmt_valor(t['valor'], es_pct, es_ratio)
                c = ws.cell(row=row, column=col, value=v)
                c.number_format = f
                c.font = NORM_FONT
                c.border = border
                col += 1

                vr, _ = fmt_valor(t['vari_rel'], True, False)
                c = ws.cell(row=row, column=col, value=vr)
                c.number_format = PCT_FMT
                c.font = NORM_FONT
                c.border = border
                col += 1

            v, f = fmt_valor(ind['anio_actual'], es_pct, es_ratio)
            c = ws.cell(row=row, column=col, value=v)
            c.number_format = f
            c.font = BOLD_FONT
            c.border = border
            col += 1

            vr, _ = fmt_valor(ind['vari_rel_aa'], True, False)
            c = ws.cell(row=row, column=col, value=vr)
            c.number_format = PCT_FMT
            c.font = NORM_FONT
            c.border = border

            row += 1

        ws.freeze_panes = 'B3'

        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(path)
        return path
