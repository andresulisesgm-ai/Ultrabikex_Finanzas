from openpyxl.styles import Font, Alignment

MONTH_TYPE = {
    'ENE': 'A', 'FEB': 'B', 'MAR': 'C', 'ABR': 'B', 'MAY': 'B', 'JUN': 'D',
    'JUL': 'B', 'AGO': 'B', 'SEPT': 'C', 'OCT': 'B', 'NOV': 'B', 'DIC': 'E'
}

HEADER_FONT = Font(name='Arial', bold=True)
HEADER_ALIGN = Alignment(horizontal='center', vertical='center')


def _month_subcols(tipo):
    cols = ['Monto', '%V', '%G']
    if tipo == 'A':
        return cols
    cols.append('Vari Rel.')
    if tipo == 'E':
        cols += ['AÑO', '%V', '%G']
    else:
        cols += ['ACUM EJEC', '%V', '%G']
    if tipo in ('C', 'D', 'E'):
        cols += ['ACUM PPTO', '%V', 'Var PPTO']
        if tipo == 'D':
            cols += ['PROM 6 EJEC', '%V', '%G', 'PROM 6 PPTO', '%V', 'Var PPTO']
    return cols


def _blank(v):
    """Replica fmtZ/pctZ: 0 o None se muestra vacío."""
    return None if (v is None or v == 0) else v


def _month_values(mes_data, tipo, muestra_pct_gastos):
    ejec = mes_data.get('ejecutado', {}) or {}
    out = [
        (_blank(ejec.get('valor')), 'num'),
        (_blank(ejec.get('pct_vtas')), 'pct'),
        (_blank(ejec.get('pct_gastos')) if muestra_pct_gastos else None, 'pct'),
    ]
    if tipo == 'A':
        return out
    out.append((_blank(mes_data.get('vari_rel')), 'pct'))
    bloque = (mes_data.get('anio') if tipo == 'E' else mes_data.get('acum_ejecutado')) or {}
    out += [
        (_blank(bloque.get('valor')), 'num'),
        (_blank(bloque.get('pct_vtas')), 'pct'),
        (_blank(bloque.get('pct_gastos')) if muestra_pct_gastos else None, 'pct'),
    ]
    if tipo in ('C', 'D', 'E'):
        ppto = mes_data.get('acum_ppto', {}) or {}
        out += [
            (_blank(ppto.get('valor')), 'num'),
            (_blank(ppto.get('pct_vtas')), 'pct'),
            (_blank(mes_data.get('var_ppto')), 'pct'),
        ]
        if tipo == 'D':
            p6e = mes_data.get('prom_6_ejec', {}) or {}
            p6p = mes_data.get('prom_6_ppto', {}) or {}
            out += [
                (_blank(p6e.get('valor')), 'num'),
                (_blank(p6e.get('pct_vtas')), 'pct'),
                (_blank(p6e.get('pct_gastos')) if muestra_pct_gastos else None, 'pct'),
                (_blank(p6p.get('valor')), 'num'),
                (_blank(p6p.get('pct_vtas')), 'pct'),
                (_blank(mes_data.get('var_ppto_prom')), 'pct'),
            ]
    return out


def _build_eerr_header(ws, months, year_prev_label, start_row=1):
    row1, row2 = start_row, start_row + 1
    col = 1

    def _set(row, c, value):
        cell = ws.cell(row=row, column=c, value=value)
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        return cell

    _set(row1, col, 'PARTIDAS')
    ws.merge_cells(start_row=row1, start_column=col, end_row=row2, end_column=col)
    col += 1

    ws.merge_cells(start_row=row1, start_column=col, end_row=row1, end_column=col+2)
    _set(row1, col, str(year_prev_label))
    for i, label in enumerate(['Valor', '%V', '%G']):
        _set(row2, col+i, label)
    col += 3

    month_start_cols = {}
    for mes in months:
        tipo = MONTH_TYPE[mes]
        subcols = _month_subcols(tipo)
        n = len(subcols)
        month_start_cols[mes] = col
        ws.merge_cells(start_row=row1, start_column=col, end_row=row1, end_column=col+n-1)
        _set(row1, col, mes)
        for i, label in enumerate(subcols):
            _set(row2, col+i, label)
        col += n

    return col - 1, month_start_cols


def _apply_row_grouping(ws, rows, row_num_map):
    """
    Aplica agrupación jerárquica (outline) a las filas del Excel según el
    campo 'indent' de cada partida, igual que la jerarquía en pantalla.
    """
    ws.sheet_properties.outlinePr.summaryBelow = False
    for r in rows:
        partida = r['partida']
        indent = r.get('indent', 0)
        if indent <= 0:
            continue
        row_num = row_num_map.get(partida)
        if row_num is None:
            continue
        ws.row_dimensions[row_num].outlineLevel = min(indent, 7)
    ws.sheet_view.showOutlineSymbols = True


# Bloque de ingresos no operativos: en EERR_STRUCTURE está posicionado lejos
# del resto de ingresos (cerca de la sección de gastos). Para la hoja de Notas
# se reordena visualmente aquí, sin modificar EERR_STRUCTURE ni el motor de
# cálculo — ver Problema B, sesión jul-2026, validado contra hoja de Yocelin
# "N EERR RODEO".
BLOQUE_INGRESOS_NO_OPERATIVOS = [
    'Otros Ingresos no Operacionales',
    'Ingresos por alquileres',
    'Ingresos por intereses',
    'Ingresos por comisiones',
    'Ingresos por servicios administrativos',
    'Sobrante en ventas',
    'Sobrante de inventarios',
    'Ganancia en venta de activos',
    'Ganancia por tasa cambiaria',
    'Ganancia por diferencias en pagos',
]

# Partida después de la cual se inserta el bloque anterior.
ANCLA_INSERCION_INGRESOS_NO_OPERATIVOS = 'Ingresos por taller'

BLOQUE_INVENTARIOS_ORDEN = [
    'Faltante de inventarios',
    'Deterioro de inventarios',
]
ANCLA_INSERCION_INVENTARIOS = 'Multas'

PARTIDAS_OCULTAS_NOTAS = {
    # Encabezado — ya se muestra en fila 1 (A1) y fila 2, evita duplicado como fila de datos
    'ESTADO DE RESULTADOS',
    'PARTIDAS',
    # Ingresos — subtotales sin equivalente en hoja de Yocelin
    'Subtotal Ingresos por Venta de Mercancia',
    'Subtotal Ingresos por Servicios',
    'Subtotal Ingresos por Eventos',
    'Subtotal Ingresos por Taller',
    # Costo de Ventas / Utilidad Bruta — subtotales sin equivalente
    'Subtotal Costo de Ventas por Servicios',
    'Subtotal Costo de Ventas por Eventos',
    'Utilidad Bruta por Venta de Mercancia y Taller',
    'Utilidad Bruta por Servicios',
    'Utilidad Bruta por Eventos',
    # Gastos — headers de agrupación sin equivalente
    'Mantenimiento y reparaciones',
    'Viáticos administrativos',
    'Gastos de seguro',
    'Gastos de impuestos, tasas y contribuciones',
    'Depreciaciones, deterioro y Amortización',
    'Gastos Bancarios',
    'Gastos de sueldos y salarios empleados y directivos',
    'Gastos de complementos empleados y directivos',
    'Gastos de personal externo',
    'Gastos de pasivos laborales vacaciones',
    'Gastos de pasivos laborales utilidades',
    'Gastos de pasivos laborales prestaciones e intereses',
    'Gastos de pasivos laborales aportes',
    'Gastos de pasivos laborales bono de guardería',
    'Gastos de pasivos laborales HCM',
    'Gastos de salud y seguridad laboral dotación',
    'Gastos de salud y seguridad laboral fiestas y agasajos',
    'Otros gastos de personal',
    'Gastos de viáticos comerciales',
    'Gastos de fletes y envios no asociados al costo',
    'Otros gastos no asociados al costo',
    'Gastos por combustible',
    'Gastos de Stand y/o ferias comerciales',
    'Otros gastos de publicidad y promoción',
    'Gastos de patrocinio y donación',
    'Gastos de viáticos por eventos',
    'Gastos de página web',
    'Gastos de desarrollo',
    # Comisiones/Rentabilidad — sin equivalente en hoja de Yocelin
    'Utilidad antes de Comisiones por Ventas',
    'Utilidad después de Comisiones por Ventas',
    'Utilidad antes de intereses, impuestos, depreciación y amortización (EBITDA)',
    'Utilidad antes de Intereses e Impuestos (EBIT)',
    # Comisiones — partida fantasma sin dato real (ver mapping_groups_v2 id 210)
    'Gastos de comisiones por ventas taller',
    # Inventarios — header sin relación jerárquica real, siempre en cero
    'Faltante y deterioro de inventarios',
}


def _reordenar_estructura_para_notas(estructura):
    """
    Devuelve una copia de EERR_STRUCTURE con los bloques definidos abajo
    reinsertados inmediatamente después de su ancla correspondiente.
    No modifica la estructura original. Uso exclusivo de presentación en
    exportables — el motor de cálculo sigue usando EERR_STRUCTURE tal cual.
    """
    bloques = [
        (BLOQUE_INGRESOS_NO_OPERATIVOS, ANCLA_INSERCION_INGRESOS_NO_OPERATIVOS),
        (BLOQUE_INVENTARIOS_ORDEN, ANCLA_INSERCION_INVENTARIOS),
    ]

    todos_los_nombres_movidos = set()
    for bloque, _ in bloques:
        todos_los_nombres_movidos.update(bloque)

    resto = [item for item in estructura if item[0] not in todos_los_nombres_movidos]

    bloques_ordenados_por_ancla = {}
    for bloque, ancla in bloques:
        bloque_tuplas = [item for item in estructura if item[0] in set(bloque)]
        bloque_ordenado = sorted(
            bloque_tuplas,
            key=lambda item: bloque.index(item[0])
        )
        bloques_ordenados_por_ancla.setdefault(ancla, []).extend(bloque_ordenado)

    resultado = []
    for item in resto:
        resultado.append(item)
        if item[0] in bloques_ordenados_por_ancla:
            resultado.extend(bloques_ordenados_por_ancla[item[0]])
    return resultado


class ExcelExporter:
    """Genera Excel del EERR usando eerr_completo_v2_ui_adapter() — subtotales como valores del engine."""

    def __init__(self, year, units, months):
        self.year   = year
        self.units  = units
        self.months = months

    def generate(self):
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from app import eerr_completo_v2_ui_adapter
        import os, tempfile

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        HDR_FILL  = PatternFill('solid', start_color='1F3864')  # Azul oscuro - título
        SUB_FILL  = PatternFill('solid', start_color='D6DCE4')  # Header columnas
        SEC_FILL  = PatternFill('solid', start_color='F2F2F2')  # Gris suave - subtotales
        WHITE_FILL= PatternFill('solid', start_color='FFFFFF')  # Blanco - partidas
        ACUM_FILL = PatternFill('solid', start_color='FFFFFF')  # Blanco - acumulados
        VAR_FILL  = PatternFill('solid', start_color='FFFFFF')  # Blanco - variaciones
        WHITE_FONT= Font(name='Arial', bold=True, color='FFFFFF', size=10)  # Blanco sobre azul
        DARK_FONT = Font(name='Arial', bold=True, color='000000', size=10)  # Negro sobre gris
        NORM_FONT = Font(name='Arial', size=9)
        TEAL_FILL = PatternFill('solid', start_color='D6DCE4')  # Gris claro - totales principales
        NUM_FMT   = '#,##0.00;(#,##0.00);"-"'
        PCT_FMT   = '0.0%;(0.0%);"-"'
        thin      = Side(style='thin', color='D9D9D9')  # Borde gris suave
        border    = Border(left=thin, right=thin, top=thin, bottom=thin)

        TOTALES_HDR = {
            'Total Ingresos', 'Total Costo de Ventas', 'Utilidad Bruta',
            'Total Gastos Operacionales', 'Total Gastos Operacionales y No Operacionales',
            'Utilidad Neta', 'Utilidad Neta despues de ISLR',
            'Utilidad antes de intereses, impuestos, depreciación y amortización (EBITDA)',
            'Utilidad antes de Intereses e Impuestos (EBIT)',
            'Utilidad Bruta por Venta de Mercancia y Taller',
            'Utilidad Bruta por Servicios', 'Utilidad Bruta por Eventos',
            'Utilidad antes de Comisiones por Ventas',
            'Utilidad después de Comisiones por Ventas'
        }

        sheets_to_build = list(self.units) + ['CONSOLIDADO']

        for unit in sheets_to_build:
            engine_unit = unit if unit != 'CONSOLIDADO' else ''
            data = eerr_completo_v2_ui_adapter(self.year, engine_unit)
            rows = data.get('rows', [])

            ws = wb.create_sheet(unit)
            ws.sheet_view.showGridLines = False

            year_prev_label = str(int(self.year) - 1)
            last_col, month_start_cols = _build_eerr_header(ws, self.months, year_prev_label, start_row=2)

            ws.merge_cells(f'A1:{get_column_letter(last_col)}1')
            title = ws['A1']
            title.value     = f'ESTADO DE RESULTADOS — {unit.upper()} — {self.year}'
            title.font      = Font(name='Arial', bold=True, color='FFFFFF', size=12)
            title.fill      = HDR_FILL
            title.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 24
            ws.row_dimensions[2].height = 20
            ws.row_dimensions[3].height = 20

            ws.column_dimensions['A'].width = 52
            for ci in range(2, 5):
                ws.column_dimensions[get_column_letter(ci)].width = 11
            for mes, start_c in month_start_cols.items():
                n = len(_month_subcols(MONTH_TYPE[mes]))
                for i in range(n):
                    ws.column_dimensions[get_column_letter(start_c + i)].width = 11

            row_num = 4
            partida_rows = {}
            for item in rows:
                partida   = item['partida']
                is_header = item['is_header']
                indent    = item.get('indent', 0)
                meses     = item.get('meses', [])
                year_prev = item.get('year_prev', {}) or {}
                muestra_pct_gastos = item.get('muestra_pct_gastos', False)

                ws.row_dimensions[row_num].height = 16

                if partida in TOTALES_HDR:
                    fill, fnt = TEAL_FILL, DARK_FONT
                elif is_header:
                    fill, fnt = SEC_FILL, DARK_FONT
                else:
                    fill, fnt = WHITE_FILL, NORM_FONT

                a_cell = ws.cell(row=row_num, column=1, value=partida)
                a_cell.font      = fnt
                a_cell.fill      = fill
                a_cell.alignment = Alignment(vertical='center', indent=indent)
                a_cell.border    = border

                yp_vals = [
                    (_blank(year_prev.get('valor')), 'num'),
                    (_blank(year_prev.get('pct_vtas')), 'pct'),
                    (_blank(year_prev.get('pct_gastos')) if muestra_pct_gastos else None, 'pct'),
                ]
                for i, (val, fmt) in enumerate(yp_vals):
                    if fmt == 'pct' and val is not None:
                        val = val / 100
                    c = ws.cell(row=row_num, column=2 + i, value=val)
                    c.number_format = PCT_FMT if fmt == 'pct' else NUM_FMT
                    c.font      = fnt
                    c.fill      = fill
                    c.alignment = Alignment(horizontal='right', vertical='center')
                    c.border    = border

                for month in self.months:
                    mes_data = next((m for m in meses if m.get('month') == month), {}) or {}
                    tipo = MONTH_TYPE[month]
                    start_c = month_start_cols[month]
                    for i, (val, fmt) in enumerate(_month_values(mes_data, tipo, muestra_pct_gastos)):
                        if fmt == 'pct' and val is not None:
                            val = val / 100
                        c = ws.cell(row=row_num, column=start_c + i, value=val)
                        c.number_format = PCT_FMT if fmt == 'pct' else NUM_FMT
                        c.font      = fnt
                        c.fill      = fill
                        c.alignment = Alignment(horizontal='right', vertical='center')
                        c.border    = border

                partida_rows[partida] = row_num
                row_num += 1

            ws.freeze_panes = f'{get_column_letter(month_start_cols[self.months[0]])}4'

            notes_name, partida_month_rows = self._build_notes_sheet(wb, unit, engine_unit)

            if partida_month_rows:
                months_list = list(self.months)
                for partida, r in partida_rows.items():
                    if partida in partida_month_rows:
                        for month in months_list:
                            filas = partida_month_rows[partida].get(month, [])
                            if filas:
                                refs = '+'.join(
                                    f"'{notes_name}'!{get_column_letter(3 + months_list.index(month))}{f}"
                                    for f in filas
                                )
                                col_monto = month_start_cols[month]
                                c = ws.cell(row=r, column=col_monto)
                                c.value = f'={refs}'
                                c.number_format = NUM_FMT

            _apply_row_grouping(ws, rows, partida_rows)

        path = os.path.join(tempfile.gettempdir(), f'EEFF_ULTRAX_{self.year}.xlsx')
        wb.save(path)
        return path

    def _build_notes_sheet(self, wb, unit, engine_unit):
        import os
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from collections import defaultdict
        from app import eerr_completo_v2_ui_adapter
        from engine import EERR_STRUCTURE

        HDR_FILL  = PatternFill('solid', start_color='1F3864')
        SEC_FILL  = PatternFill('solid', start_color='F2F2F2')
        WHITE_FILL= PatternFill('solid', start_color='FFFFFF')
        WHITE_FONT= Font(name='Arial', bold=True, color='FFFFFF', size=10)
        DARK_FONT = Font(name='Arial', bold=True, color='000000', size=10)
        NORM_FONT = Font(name='Arial', size=9)
        NUM_FMT   = '#,##0.00;(#,##0.00);"-"'
        thin      = Side(style='thin', color='D9D9D9')
        border    = Border(left=thin, right=thin, top=thin, bottom=thin)

        notes_sheet_name = f'N EERR {unit}'
        ws = wb.create_sheet(notes_sheet_name)
        ws.sheet_view.showGridLines = False

        data_adapter = eerr_completo_v2_ui_adapter(self.year, engine_unit)
        rows_by_partida = {r['partida']: r for r in data_adapter.get('rows', [])}

        N = len(self.months)
        total_cols = 2 + N + 1  # partida, year_prev, meses..., total

        # Columna A: "ESTADO DE RESULTADOS" suelto, replicando layout de Yocelin.
        a1 = ws['A1']
        a1.value = 'ESTADO DE RESULTADOS'
        a1.font = Font(name='Arial', bold=True, color='FFFFFF', size=12)
        a1.fill = HDR_FILL
        a1.alignment = Alignment(horizontal='center', vertical='center')

        # Título mergeado desde columna B (ya no ocupa la columna A).
        ws.merge_cells(f'B1:{get_column_letter(total_cols)}1')
        t = ws['B1']
        t.value = f'NOTAS — ESTADO DE RESULTADOS — {unit.upper()} — {self.year}'
        t.font = Font(name='Arial', bold=True, color='FFFFFF', size=12)
        t.fill = HDR_FILL
        t.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 24

        headers = ['PARTIDAS', str(data_adapter.get('year_prev', int(self.year) - 1))] + list(self.months) + ['TOTAL']
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=2, column=ci, value=h)
            cell.font = WHITE_FONT
            cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        ws.column_dimensions['A'].width = 45
        for ci in range(2, total_cols + 1):
            ws.column_dimensions[get_column_letter(ci)].width = 12

        partida_month_rows = defaultdict(lambda: defaultdict(list))
        row_num = 3

        def _escribir_fila_notas(nombre_mostrado, r, es_header, row_num):
            ws.cell(row=row_num, column=1, value=nombre_mostrado).border = border
            ws.cell(row=row_num, column=1).font = DARK_FONT if es_header else NORM_FONT

            yp = r.get('year_prev', {}).get('valor', 0)
            ws.cell(row=row_num, column=2, value=yp).number_format = NUM_FMT
            ws.cell(row=row_num, column=2).border = border

            for m_idx, month in enumerate(self.months):
                mes_data = next((m for m in r.get('meses', []) if m.get('month') == month), {})
                val = mes_data.get('ejecutado', {}).get('valor', 0)
                col = 3 + m_idx
                c = ws.cell(row=row_num, column=col, value=val)
                c.number_format = NUM_FMT
                c.font = NORM_FONT
                c.border = border
                c.alignment = Alignment(horizontal='right')

                if not es_header:
                    partida_month_rows[nombre_mostrado][month].append(row_num)

            m_start = get_column_letter(3)
            m_end = get_column_letter(2 + N)
            tot = ws.cell(row=row_num, column=3 + N,
                          value=f'=SUM({m_start}{row_num}:{m_end}{row_num})')
            tot.number_format = NUM_FMT
            tot.font = NORM_FONT
            tot.border = border
            tot.alignment = Alignment(horizontal='right')

        for item in _reordenar_estructura_para_notas(EERR_STRUCTURE):
            partida_nombre = item[0]
            es_header = item[1]

            # Se omite aquí: se muestra como "Total Gastos" justo después de
            # "Utilidad Bruta", replicando la posición de la hoja de Yocelin.
            if partida_nombre == 'Total Gastos Operacionales y No Operacionales':
                continue

            if partida_nombre in PARTIDAS_OCULTAS_NOTAS:
                continue

            r = rows_by_partida.get(partida_nombre)
            if r is None:
                continue

            _escribir_fila_notas(partida_nombre, r, es_header, row_num)
            row_num += 1

            # Fila sintética "Total Gastos" — solo presentación en el exportable,
            # no existe en EERR_STRUCTURE ni en el motor de cálculo. Reutiliza el
            # valor ya calculado por el motor bajo 'Total Gastos Operacionales y
            # No Operacionales' (= gastos_operacionales + otros_gastos).
            if partida_nombre == 'Utilidad Bruta':
                r_total_gastos = rows_by_partida.get('Total Gastos Operacionales y No Operacionales')
                if r_total_gastos is not None:
                    _escribir_fila_notas('Total Gastos', r_total_gastos, True, row_num)
                    row_num += 1

        ws.freeze_panes = 'A3'
        return notes_sheet_name, partida_month_rows
