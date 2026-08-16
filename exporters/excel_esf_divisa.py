NUM_FMT = '#,##0.00;-#,##0.00;"-"'
PCT_FMT = '0.0%;(0.0%);"-"'

EXCLUIR_ESF_EXPORT = {
    'Ajuste por Diferencial Cambiario', 'Propiedades de inversión',
    'Total Activos Corrientes', 'Total Activos No Corrientes',
    'Total Pasivos Corrientes', 'Total Pasivos No Corrientes',
    'Total Patrimonio'
}
RENOMBRAR_ESF_EXPORT = {
    # Activos Corrientes
    'Caja en Bs': 'Cajas en Bs',
    'Caja en $': 'Cajas en $',
    'Fondo en Bs': 'Fondos en Bs',
    'Fondo en $': 'Fondos en $',
    'Bancos en Bs': 'Bancos en Bs',
    'Bancos en $': 'Bancos en $',
    'Bancos en transito en Bs': 'Bancos en transito en Bs',
    'Bancos en transito en $': 'Bancos en transito en $',
    'Cuentas por cobrar clientes': 'A clientes',
    'A empresas relacionadas del grupo (CxC)': 'A empresas relacionadas del grupo',
    'A empresas externas del grupo (CxC)': 'A empresas externas del grupo',
    'A socios (CxC)': 'A socios',
    'Cuentas por cobrar empleados': 'A empleados',
    'Cuentas por cobrar por sociedades': 'Por sociedades',
    'A empresas relacionadas del grupo (PxC)': 'A empresas relacionadas del grupo',
    'A empresas externas del grupo (PxC)': 'A empresas externas del grupo',
    'A socios (PxC)': 'A socios',
    'A empleados (PxC)': 'A empleados',
    'Otros prestamos por cobrar': 'A terceros',
    'Anticipos a proveedores': 'A proveedores',
    'Anticipos a socios': 'A socios',
    'Anticipos a empleados': 'A empleados',
    'Inventario de mercancias': 'De mercancía',
    'Inventario de suministros': 'De suministros',
    'Inventario en transito': 'En tránsito',
    'Inventarios de Consignacion': 'A consignación',
    'Impuestos pagados por anticipado': 'Impuestos ',

    # Activos No Corrientes
    'Cuentas por cobrar clientes L.P.': 'Deudores comerciales a LP',
    'Otras cuentas por cobrar L.P.': 'Otras cuentas por cobrar a LP',
    'Prestamos por cobrar L.P.': 'Préstamos por cobrar LP',
    'Propiedades de Inversión': 'Propiedades de inversión',

    # Pasivos Corrientes
    'A proveedores': 'A proveedores ',
    'A empresas relacionadas del grupo (CxP)': 'A empresas relacionadas del grupo',
    'A empresas externas del grupo (CxP)': 'A empresas externas del grupo',
    'A socios (CxP)': 'A socios',
    'Otras Cuentas por pagar': 'Otras cuentas por pagar',
    'Retenciones Laborales por pagar': 'Retenciones laborales a pagar',
    'A empresas relacionadas del grupo (PxP)': 'A empresas relacionadas del grupo',
    'A empresas externas del grupo (PxP)': 'A empresas externas del grupo',
    'A socios (PxP)': 'A socios',
    'A empleados (PxP)': 'A empleados',
    'Otros prestamos por pagar': 'Otros prestamos por pagar ',
    'Anticipos de Pasivo': 'Anticipos',
}
REPARENT_ESF_EXPORT = {
    'Propiedades de Inversión': ('Otros activos no corrientes', 3),
}

NOMBRES_ESPECIALES = {
    'Resultados del ejercicio', 'Resultados acumulados',
    'ACTIVOS CORRIENTES', 'ACTIVOS NO CORRIENTES',
    'PASIVOS CORRIENTES', 'PASIVOS NO CORRIENTES', 'PATRIMONIO',
    'Total Activos Corrientes', 'Total Activos No Corrientes',
    'Total Pasivos Corrientes', 'Total Pasivos No Corrientes',
    'Total Patrimonio',
    'TOTAL ACTIVOS', 'TOTAL PASIVOS', 'TOTAL PASIVOS Y PATRIMONIO',
}


class ESFDivisaRealExporter:
    """Genera ESF ULTRAX DIVISA REAL consolidado con fórmulas vivas, aplicando las reglas y
    overrides de Divisa Real a través de esf_engine() y calcular_esf_divisa_real()."""

    def __init__(self, year, empresa_id=None):
        from engine import MONTHS
        self.year = year
        self.empresa_id = empresa_id
        self.months = list(MONTHS)

    def generate(self):
        import openpyxl, os, tempfile
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from engine import esf_engine, ESF_STRUCTURE_V2, calcular_esf_divisa_real, PARTIDAS_ESF_DIVISA_REAL
        from exporters.excel_eerr import ExcelExporter
        from db import get_db

        year_prev = str(int(self.year) - 1)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        db_conn = get_db()

        from helpers import _nombre_empresa_display
        nombre_emp = _nombre_empresa_display(db_conn, self.empresa_id)

        # Unidades reales de la empresa activa (o la lista de todas si es Holding/empresa_id=None).
        if self.empresa_id is not None:
            units_export = [r[0] for r in db_conn.execute('SELECT nombre FROM unidades WHERE empresa_id=?', [self.empresa_id]).fetchall()]
        else:
            units_export = [r[0] for r in db_conn.execute('SELECT nombre FROM unidades').fetchall()]

        # Overrides de divisa real y paneles por trimestre
        overrides_combinado = {}
        self.panel_por_quarter = {}
        self.quarters_sin_tasa = []
        for q in (1, 2, 3, 4):
            overrides_rows = db_conn.execute(
                'SELECT odoo_code, valor_override FROM esf_divisa_real_overrides WHERE year=? AND quarter=? AND empresa_id IS ?',
                (self.year, q, self.empresa_id)
            ).fetchall()
            for r in overrides_rows:
                overrides_combinado[(q, r['odoo_code'])] = r['valor_override']

            panel_q = calcular_esf_divisa_real(self.year, q, empresa_id=self.empresa_id)
            if 'error' in panel_q:
                self.quarters_sin_tasa.append(q)
            else:
                self.panel_por_quarter[q] = panel_q

        # 1. EERR ULTRAX (año actual): unidades reales + notas, reusando el exportador existente.
        eerr_meta = ExcelExporter(self.year, units_export, self.months, empresa_id=self.empresa_id).generate(
            wb=wb, save=False, return_meta=True
        )
        # 2. EERR ULTRAX (Año Anterior): solo consolidado, sin unidades ni notas.
        eerr_prev_meta = ExcelExporter(year_prev, [], self.months, empresa_id=self.empresa_id).generate(
            wb=wb, save=False, return_meta=True, consolidado_name=f'EERR {nombre_emp} (Año Ant.)'[:31]
        )

        sheet_curr = eerr_meta['sheet_name_consolidado']
        row_islr_curr = eerr_meta['partida_rows_consolidado'].get('Utilidad Neta despues de ISLR')
        cols_curr = eerr_meta['month_start_cols']

        sheet_prev = eerr_prev_meta['sheet_name_consolidado']
        row_islr_prev = eerr_prev_meta['partida_rows_consolidado'].get('Utilidad Neta despues de ISLR')
        cols_prev = eerr_prev_meta['month_start_cols']

        # 3. Datos ESF vía orquestador de metodología EERR Real / ESF Real (ago-2026):
        # Acumulados fijo (=BCV), diferencia mandada como plug al EERR Real -- ver
        # calcular_estados_reales en engine.py.
        from engine import calcular_estados_reales
        estados_reales = calcular_estados_reales(self.year, '', empresa_id=self.empresa_id)
        data_curr = estados_reales['esf_real']
        data_prev = esf_engine(year_prev, '', empresa_id=self.empresa_id)

        expected_levels = {item[0]: item[5] for item in ESF_STRUCTURE_V2}
        qdata_curr = {}
        for r in data_curr['rows']:
            partida = r['partida']
            if partida in expected_levels and r.get('level') == expected_levels[partida]:
                qdata_curr[partida] = r['quarters']

        qdata_prev = {}
        for r in data_prev['rows']:
            partida = r['partida']
            if partida in expected_levels and r.get('level') == expected_levels[partida]:
                qdata_prev[partida] = r['quarters']

        # 4. Estructura filtrada/ajustada para el exportable.
        filtered = []
        for item in ESF_STRUCTURE_V2:
            name, is_header, _, bold, bg_color, level, _, parent_name, indent = item
            if name in EXCLUIR_ESF_EXPORT:
                continue
            if name in REPARENT_ESF_EXPORT:
                parent_name, level = REPARENT_ESF_EXPORT[name]
                is_header = False
            display = RENOMBRAR_ESF_EXPORT.get(name, name)
            filtered.append({
                'name': name, 'display': display, 'is_header': is_header,
                'bold': bold, 'level': level, 'parent_name': parent_name, 'indent': indent,
            })

        # Reordenamiento de TOTAL ACTIVOS para moverlo al final de los activos
        idx_ta = None
        for i, f in enumerate(filtered):
            if f['name'] == 'TOTAL ACTIVOS':
                idx_ta = i
                break
        if idx_ta is not None:
            ta_item = filtered.pop(idx_ta)
            idx_pc = None
            for i, f in enumerate(filtered):
                if f['name'] == 'PASIVOS CORRIENTES':
                    idx_pc = i
                    break
            if idx_pc is not None:
                filtered.insert(idx_pc, ta_item)
            else:
                filtered.append(ta_item)

        HDR_FILL  = PatternFill('solid', start_color='1F3864')
        SEC_FILL  = PatternFill('solid', start_color='F2F2F2')
        WHITE_FILL= PatternFill('solid', start_color='FFFFFF')
        TEAL_FILL = PatternFill('solid', start_color='D6DCE4')
        DARK_FONT = Font(name='Arial', bold=True, color='000000', size=10)
        NORM_FONT = Font(name='Arial', size=9)
        thin      = Side(style='thin', color='D9D9D9')
        border    = Border(left=thin, right=thin, top=thin, bottom=thin)

        nombre_hoja_esf_dr = f'ESF {nombre_emp} DIVISA REAL'[:31]
        ws = wb.create_sheet(nombre_hoja_esf_dr)
        ws.sheet_view.showGridLines = False

        headers = ['PARTIDA', 'Año Anterior', 'Q1', 'Vari Rel.', 'Q2', 'Vari Rel.', 'Q3', 'Vari Rel.', 'Q4', 'Vari Rel.', 'Vari Rel. actual vs anterior']
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=2, column=ci, value=h)
            c.font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
            c.fill = HDR_FILL
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = border

        ws.merge_cells('A1:K1')
        title = ws['A1']
        title.value = f'ESTADO DE SITUACIÓN FINANCIERA — DIVISA REAL — {nombre_emp.upper()} — {self.year}'
        title.font = Font(name='Arial', bold=True, color='FFFFFF', size=12)
        title.fill = HDR_FILL
        title.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 24

        ws.column_dimensions['A'].width = 48
        for ci in range(2, 12):
            ws.column_dimensions[get_column_letter(ci)].width = 14

        COL_AA, COL_Q1, COL_V1, COL_Q2, COL_V2, COL_Q3, COL_V3, COL_Q4, COL_V4, COL_VF = range(2, 12)

        row_num = 3
        row_of = {}
        for f in filtered:
            fill, fnt = (TEAL_FILL, DARK_FONT) if f['name'] in NOMBRES_ESPECIALES else \
                        (SEC_FILL, DARK_FONT) if f['is_header'] else (WHITE_FILL, NORM_FONT)

            a = ws.cell(row=row_num, column=1, value=f['display'])
            a.font = fnt
            a.fill = fill
            a.alignment = Alignment(vertical='center', indent=f['indent'])
            a.border = border

            qc = qdata_curr.get(f['name'], {1: 0.0, 2: 0.0, 3: 0.0, 4: 0.0})
            qp = qdata_prev.get(f['name'], {1: 0.0, 2: 0.0, 3: 0.0, 4: 0.0})

            valores = [qp.get(4, 0.0), qc.get(1, 0.0), None, qc.get(2, 0.0), None,
                       qc.get(3, 0.0), None, qc.get(4, 0.0), None, None]
            for i, val in enumerate(valores):
                if val is None:
                    continue
                c = ws.cell(row=row_num, column=2 + i, value=val)
                c.number_format = NUM_FMT
                c.font = fnt
                c.fill = fill
                c.border = border
                c.alignment = Alignment(horizontal='right', vertical='center')

            row_of[f['name']] = row_num
            row_num += 1

        def L(col):
            return get_column_letter(col)

        def set_formula(name, col, formula):
            r = row_of.get(name)
            if r is None:
                return
            c = ws.cell(row=r, column=col, value=formula)
            c.number_format = NUM_FMT

        # --- Nivel 2: headers que suman sus hijos nivel 3 ---
        hijos_por_padre = {}
        for f in filtered:
            hijos_por_padre.setdefault(f['parent_name'], []).append(f['name'])

        headers_nivel2 = [f['name'] for f in filtered if f['is_header'] and f['name'] not in NOMBRES_ESPECIALES]
        for padre in headers_nivel2:
            hijos = hijos_por_padre.get(padre, [])
            hijos_rows = [row_of[h] for h in hijos if h in row_of]
            if not hijos_rows:
                continue
            for col in (COL_AA, COL_Q1, COL_Q2, COL_Q3, COL_Q4):
                refs = '+'.join(f'{L(col)}{hr}' for hr in hijos_rows)
                set_formula(padre, col, f'={refs}')

        # --- CATEGORIAS_NIVEL1: Se calculan sumando sus hijos nivel 2 directos ---
        CATEGORIAS_NIVEL1 = {
            'ACTIVOS CORRIENTES': ['Efectivo y Equivalentes', 'Cuentas por Cobrar', 'Otras Cuentas por Cobrar', 'Préstamos por Cobrar', 'Anticipos', 'Inventarios', 'Prepagados'],
            'ACTIVOS NO CORRIENTES': ['Otros activos no corrientes', 'Propiedades, Plantas y Equipos'],
            'PASIVOS CORRIENTES': ['Cuentas por Pagar', 'Otras cuentas por pagar', 'Préstamos por Pagar', 'Anticipos de Pasivo', 'Provisiones'],
            'PASIVOS NO CORRIENTES': ['Otras cuentas por pagar L.P.'],
            'PATRIMONIO': ['Capital social', 'Reservas legales y estatutarias', 'Superavit por revaluacion', 'Resultados acumulados', 'Resultados del ejercicio'],
        }

        for cat, hijos in CATEGORIAS_NIVEL1.items():
            hijos_rows = [row_of[h] for h in hijos if h in row_of]
            if not hijos_rows:
                continue
            for col in (COL_AA, COL_Q1, COL_Q2, COL_Q3, COL_Q4):
                refs = '+'.join(f'{L(col)}{hr}' for hr in hijos_rows)
                set_formula(cat, col, f'={refs}')

        # --- Nivel 0: TOTAL ACTIVOS / TOTAL PASIVOS / TOTAL PASIVOS Y PATRIMONIO ---
        r_ac = row_of.get('ACTIVOS CORRIENTES')
        r_anc = row_of.get('ACTIVOS NO CORRIENTES')
        r_pc = row_of.get('PASIVOS CORRIENTES')
        r_pnc = row_of.get('PASIVOS NO CORRIENTES')
        r_pat = row_of.get('PATRIMONIO')
        r_ta = row_of.get('TOTAL ACTIVOS')
        r_tp = row_of.get('TOTAL PASIVOS')
        for col in (COL_AA, COL_Q1, COL_Q2, COL_Q3, COL_Q4):
            if r_ac and r_anc:
                set_formula('TOTAL ACTIVOS', col, f'={L(col)}{r_ac}+{L(col)}{r_anc}')
            if r_pc and r_pnc:
                set_formula('TOTAL PASIVOS', col, f'={L(col)}{r_pc}+{L(col)}{r_pnc}')
            if r_tp and r_pat:
                set_formula('TOTAL PASIVOS Y PATRIMONIO', col, f'={L(col)}{r_tp}+{L(col)}{r_pat}')

        # --- Resultados del ejercicio ---
        ultimo_mes_q = {
            1: 'MAR',
            2: 'JUN',
            3: 'SEPT',
            4: self.months[-1] if self.months else 'DIC',
        }
        # Metodología EERR Real / ESF Real (ago-2026): "Resultados del ejercicio" del
        # año actual ya NO referencia la hoja EERR BCV embebida -- el valor correcto
        # (Utilidad Neta real preliminar, sin la línea cambiaria) ya quedó escrito
        # como literal en el loop de población inicial, desde data_curr.
        if row_islr_prev:
            meses_prev = [m for m in self.months if m in cols_prev]
            if meses_prev:
                ultimo_mes_prev = meses_prev[-1]
                col_acum_prev = cols_prev[ultimo_mes_prev] + 4
                ref = f"'{sheet_prev}'!{L(col_acum_prev)}{row_islr_prev}"
                set_formula('Resultados del ejercicio', COL_AA, f'={ref}')

        # Metodología EERR Real / ESF Real (ago-2026): "Resultados acumulados" ya NO
        # se recalcula como plug/cuadre en Excel -- queda fijo, idéntico al de BCV,
        # tal como exige la regla de Yocelin. El valor ya está escrito como literal
        # en el loop de población inicial, desde data_curr (esf_real).

        # --- Vari Rel. ---
        VARI_PARES = [
            (COL_V1, COL_Q1, COL_AA),
            (COL_V2, COL_Q2, COL_Q1),
            (COL_V3, COL_Q3, COL_Q2),
            (COL_V4, COL_Q4, COL_Q3),
            (COL_VF, COL_Q4, COL_AA),
        ]
        for f in filtered:
            r = row_of[f['name']]
            for col_v, col_actual, col_base in VARI_PARES:
                actual_ref = f'{L(col_actual)}{r}'
                base_ref = f'{L(col_base)}{r}'
                formula = (
                    f'=IF(AND({base_ref}=0,{actual_ref}>0),100%,'
                    f'IF(AND({base_ref}<0,{actual_ref}>=0),({base_ref}-{actual_ref})/{base_ref},'
                    f'IF(AND({base_ref}>0,{actual_ref}>0),({actual_ref}-{base_ref})/{base_ref},'
                    f'IF(AND({base_ref}=0,{actual_ref}<=0),"",'
                    f'({actual_ref}-{base_ref})/{base_ref}))))'
                )
                c = ws.cell(row=r, column=col_v, value=formula)
                c.number_format = PCT_FMT
                c.font = NORM_FONT
                c.border = border
                c.alignment = Alignment(horizontal='right', vertical='center')

        ws.sheet_properties.outlinePr.summaryBelow = False
        for f in filtered:
            if f.get('level') == 3:
                r = row_of.get(f['name'])
                if r is not None:
                    ws.row_dimensions[r].outlineLevel = 1
                    ws.row_dimensions[r].hidden = True
        ws.sheet_view.showOutlineSymbols = True

        ws.freeze_panes = 'B3'

        # 5. Hoja adicional "Divisa Real - Detalle" (Panel de 7 partidas)
        ws_dr = wb.create_sheet('Divisa Real - Detalle')
        ws_dr.sheet_view.showGridLines = False

        HDR_FILL_DR = PatternFill('solid', start_color='1F3864')
        NO_TASA_FILL = PatternFill('solid', start_color='D9D9D9')
        TOTAL_FILL = PatternFill('solid', start_color='F2F2F2')
        WHITE_FILL_DR = PatternFill('solid', start_color='FFFFFF')

        FONT_TITLE_DR = Font(name='Arial', bold=True, color='FFFFFF', size=12)
        FONT_HDR_DR = Font(name='Arial', bold=True, color='FFFFFF', size=10)
        FONT_BOLD_DR = Font(name='Arial', bold=True, color='000000', size=9)
        FONT_NORM_DR = Font(name='Arial', size=9)
        FONT_WARN_DR = Font(name='Arial', italic=True, color='595959', size=8)

        thin_dr = Side(style='thin', color='D9D9D9')
        border_dr = Border(left=thin_dr, right=thin_dr, top=thin_dr, bottom=thin_dr)

        ws_dr.merge_cells('A1:E1')
        t_dr = ws_dr['A1']
        t_dr.value = f'PARTIDAS REVALORIZABLES — DIVISA REAL — {nombre_emp.upper()} — {self.year}'
        t_dr.font = FONT_TITLE_DR
        t_dr.fill = HDR_FILL_DR
        t_dr.alignment = Alignment(horizontal='center', vertical='center')
        ws_dr.row_dimensions[1].height = 24

        headers_dr = ['PARTIDA', 'Q1', 'Q2', 'Q3', 'Q4']
        for ci, h in enumerate(headers_dr, 1):
            c = ws_dr.cell(row=2, column=ci, value=h)
            c.font = FONT_HDR_DR
            c.fill = HDR_FILL_DR
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = border_dr

        ws_dr.column_dimensions['A'].width = 40
        for ci in range(2, 6):
            ws_dr.column_dimensions[get_column_letter(ci)].width = 16

        current_row = 3
        # 5a. Filas de partidas revalorizables
        for idx_p, p_name in enumerate(PARTIDAS_ESF_DIVISA_REAL):
            cell_p = ws_dr.cell(row=current_row, column=1, value=p_name)
            cell_p.font = FONT_NORM_DR
            cell_p.fill = WHITE_FILL_DR
            cell_p.border = border_dr
            cell_p.alignment = Alignment(vertical='center')

            for q, col_idx in zip((1, 2, 3, 4), range(2, 6)):
                c = ws_dr.cell(row=current_row, column=col_idx)
                c.border = border_dr
                if q in self.quarters_sin_tasa:
                    c.fill = NO_TASA_FILL
                    if idx_p == 0:
                        c.value = "Sin tasa paralela cargada"
                        c.font = FONT_WARN_DR
                        c.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    c.fill = WHITE_FILL_DR
                    c.font = FONT_NORM_DR
                    c.number_format = NUM_FMT
                    c.alignment = Alignment(horizontal='right', vertical='center')
                    panel_q = self.panel_por_quarter.get(q, {})
                    partidas_dict = panel_q.get('partidas', {})
                    c.value = partidas_dict.get(p_name, 0.0)

            current_row += 1

        # 5b. Fila Tasa Paralela Fin de Mes
        cell_tasa = ws_dr.cell(row=current_row, column=1, value='Tasa Paralela Fin de Mes')
        cell_tasa.font = FONT_BOLD_DR
        cell_tasa.fill = TOTAL_FILL
        cell_tasa.border = border_dr
        cell_tasa.alignment = Alignment(vertical='center')

        for q, col_idx in zip((1, 2, 3, 4), range(2, 6)):
            c = ws_dr.cell(row=current_row, column=col_idx)
            c.border = border_dr
            if q in self.quarters_sin_tasa:
                c.fill = NO_TASA_FILL
            else:
                c.fill = TOTAL_FILL
                c.font = FONT_BOLD_DR
                c.number_format = NUM_FMT
                c.alignment = Alignment(horizontal='right', vertical='center')
                panel_q = self.panel_por_quarter.get(q, {})
                c.value = panel_q.get('tasa_paralela_fin', 0.0)
        current_row += 1

        # 5c. Fila TOTAL (Bs)
        cell_tot_bs = ws_dr.cell(row=current_row, column=1, value='TOTAL (Bs)')
        cell_tot_bs.font = FONT_BOLD_DR
        cell_tot_bs.fill = TOTAL_FILL
        cell_tot_bs.border = border_dr
        cell_tot_bs.alignment = Alignment(vertical='center')

        for q, col_idx in zip((1, 2, 3, 4), range(2, 6)):
            c = ws_dr.cell(row=current_row, column=col_idx)
            c.border = border_dr
            if q in self.quarters_sin_tasa:
                c.fill = NO_TASA_FILL
            else:
                c.fill = TOTAL_FILL
                c.font = FONT_BOLD_DR
                c.number_format = NUM_FMT
                c.alignment = Alignment(horizontal='right', vertical='center')
                panel_q = self.panel_por_quarter.get(q, {})
                c.value = panel_q.get('saldo_total_bs', 0.0)
        current_row += 1

        # 5d. Fila TOTAL (USD)
        cell_tot_usd = ws_dr.cell(row=current_row, column=1, value='TOTAL (USD)')
        cell_tot_usd.font = FONT_BOLD_DR
        cell_tot_usd.fill = TOTAL_FILL
        cell_tot_usd.border = border_dr
        cell_tot_usd.alignment = Alignment(vertical='center')

        for q, col_idx in zip((1, 2, 3, 4), range(2, 6)):
            c = ws_dr.cell(row=current_row, column=col_idx)
            c.border = border_dr
            if q in self.quarters_sin_tasa:
                c.fill = NO_TASA_FILL
            else:
                c.fill = TOTAL_FILL
                c.font = FONT_BOLD_DR
                c.number_format = NUM_FMT
                c.alignment = Alignment(horizontal='right', vertical='center')
                panel_q = self.panel_por_quarter.get(q, {})
                c.value = panel_q.get('saldo_total_usd', 0.0)
        current_row += 1

        wb.move_sheet(nombre_hoja_esf_dr, offset=-(wb.sheetnames.index(nombre_hoja_esf_dr)))
        wb.move_sheet('Divisa Real - Detalle', offset=-(wb.sheetnames.index('Divisa Real - Detalle') - 1))
        wb.move_sheet(eerr_meta['sheet_name_consolidado'], offset=-(wb.sheetnames.index(eerr_meta['sheet_name_consolidado']) - 2))

        path = os.path.join(tempfile.gettempdir(), f'ESF_DIVISA_REAL_{nombre_emp}_{self.year}.xlsx')
        wb.save(path)
        return path
